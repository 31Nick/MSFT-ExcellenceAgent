"""Parser for APRL v2 Excel reports using openpyxl + pandas."""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass, field
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from excellence_agent.ingest.schema import (
    COL_RESOURCE_TYPE,
    HEADER_ROW,
    IMPACTED_RESOURCES_COLUMNS,
    PLATFORM_ISSUES_COLUMNS,
    REQUIRED_COLUMNS,
    SHEET_IMPACTED_RESOURCES,
    SHEET_PLATFORM_ISSUES,
    validate_columns,
)

logger = logging.getLogger(__name__)


@dataclass
class APRLReport:
    """Container for the parsed APRL report."""

    impacted_resources: pd.DataFrame
    platform_issues: pd.DataFrame
    metadata: dict[str, Any] = field(default_factory=dict)


class APRLParser:
    """Read an APRL v2 Excel workbook and return clean DataFrames."""

    def __init__(self, file_path: str) -> None:
        self.file_path = file_path
        if not os.path.isfile(self.file_path):
            raise FileNotFoundError(
                f"APRL report not found: {self.file_path}"
            )

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def parse(self) -> APRLReport:
        """Parse the full APRL report and return an :class:`APRLReport`."""
        impacted = self.parse_impacted_resources()
        platform = self.parse_platform_issues()

        wb = load_workbook(self.file_path, read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        wb.close()

        metadata: dict[str, Any] = {
            "sheet_names": sheet_names,
            "impacted_resources_rows": len(impacted),
            "platform_issues_rows": len(platform),
        }
        logger.info(
            "Parsed APRL report: %d impacted-resource rows, %d platform-issue rows",
            len(impacted),
            len(platform),
        )
        return APRLReport(
            impacted_resources=impacted,
            platform_issues=platform,
            metadata=metadata,
        )

    def parse_impacted_resources(self) -> pd.DataFrame:
        """Parse the *4.ImpactedResourcesAnalysis* sheet."""
        df = self._read_sheet(
            sheet_name=SHEET_IMPACTED_RESOURCES,
            expected_columns=IMPACTED_RESOURCES_COLUMNS,
            required_columns=REQUIRED_COLUMNS,
        )

        # Drop summary / formula rows that have no Resource Type value.
        if COL_RESOURCE_TYPE in df.columns:
            before = len(df)
            df = df[df[COL_RESOURCE_TYPE].notna() & (df[COL_RESOURCE_TYPE].astype(str).str.strip() != "")]
            dropped = before - len(df)
            if dropped:
                logger.debug(
                    "Filtered %d rows with empty Resource Type", dropped
                )

        df = df.reset_index(drop=True)
        return df

    def parse_platform_issues(self) -> pd.DataFrame:
        """Parse the *5.PlatformIssuesAnalysis* sheet."""
        df = self._read_sheet(
            sheet_name=SHEET_PLATFORM_ISSUES,
            expected_columns=PLATFORM_ISSUES_COLUMNS,
            required_columns=None,
        )
        df = df.reset_index(drop=True)
        return df

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _read_sheet(
        self,
        sheet_name: str,
        expected_columns: list[str],
        required_columns: list[str] | None,
    ) -> pd.DataFrame:
        """Load a single sheet into a DataFrame with validation."""
        wb = load_workbook(self.file_path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                logger.warning(
                    "Sheet '%s' not found in workbook. Available sheets: %s",
                    sheet_name,
                    wb.sheetnames,
                )
                return pd.DataFrame()

            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()

        if len(rows) <= HEADER_ROW:
            logger.warning(
                "Sheet '%s' has fewer rows (%d) than expected header row (%d)",
                sheet_name,
                len(rows),
                HEADER_ROW,
            )
            return pd.DataFrame()

        # Extract headers and normalise whitespace.
        raw_headers = rows[HEADER_ROW]
        headers = [
            str(h).strip() if h is not None else f"_unnamed_{i}"
            for i, h in enumerate(raw_headers)
        ]

        # Validate expected columns.
        is_valid, missing = validate_columns(headers, expected_columns)
        if not is_valid:
            logger.warning(
                "Sheet '%s' is missing expected columns: %s",
                sheet_name,
                missing,
            )

        # Validate required columns (stricter).
        if required_columns:
            req_valid, req_missing = validate_columns(headers, required_columns)
            if not req_valid:
                logger.warning(
                    "Sheet '%s' is missing REQUIRED columns: %s",
                    sheet_name,
                    req_missing,
                )

        # Build DataFrame from data rows (everything after the header).
        data_rows = rows[HEADER_ROW + 1 :]
        df = pd.DataFrame(data_rows, columns=headers)
        logger.debug(
            "Read %d data rows from sheet '%s'", len(df), sheet_name
        )
        return df
